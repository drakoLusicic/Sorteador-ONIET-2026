"""Lectura de la planilla de participantes (CSV o Excel .xlsx).

La primera fila tiene los títulos de las columnas: `nombre` y `apellido`
(obligatorias) e `id` (opcional: el número de participante; si no está, se
numeran solos). Se aceptan variantes como "Nombres", "APELLIDO" o
"ID de participante", y columnas de más (se ignoran).
"""

import csv
import io
import re
import unicodedata

MAX_FILAS = 20_000
LARGO_MAXIMO = 100  # mismo largo que las columnas de la base

TITULOS = {
    "id": {"id", "id participante", "id de participante", "nro", "nro participante",
           "numero", "numero de participante", "n", "codigo"},
    "nombre": {"nombre", "nombres"},
    "apellido": {"apellido", "apellidos"},
}


class ErrorImportacion(Exception):
    """Problema con el archivo, con un mensaje para mostrarle al usuario."""


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", texto.lower()).strip()


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\s+", " ", str(valor)).strip()


def _columnas(titulos):
    """Posición de cada columna conocida según la fila de títulos."""
    posiciones = {}
    for i, titulo in enumerate(titulos):
        normal = _normalizar(titulo)
        for campo, variantes in TITULOS.items():
            if normal in variantes and campo not in posiciones:
                posiciones[campo] = i
    faltan = [c for c in ("nombre", "apellido") if c not in posiciones]
    if faltan:
        raise ErrorImportacion(
            "La primera fila tiene que tener los títulos de las columnas: nombre, apellido "
            f"e id (opcional). No se encontró: {', '.join(faltan)}."
        )
    return posiciones


def _filas_csv(contenido):
    # UTF-8 (con o sin BOM) o, si no, la codificación de Excel en Windows.
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    # Excel en castellano guarda los CSV separados por punto y coma.
    muestra = texto[:4096]
    try:
        separador = csv.Sniffer().sniff(muestra, delimiters=",;\t").delimiter
    except csv.Error:
        separador = ";" if muestra.count(";") > muestra.count(",") else ","
    return list(csv.reader(io.StringIO(texto), delimiter=separador))


def _filas_xlsx(contenido):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depende de la instalación
        raise ErrorImportacion("Para leer archivos .xlsx falta instalar openpyxl.") from exc
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:
        raise ErrorImportacion("No se pudo abrir el archivo de Excel. ¿Es un .xlsx válido?") from exc
    try:
        return [list(fila) for fila in libro.worksheets[0].iter_rows(values_only=True)]
    finally:
        libro.close()


def leer_participantes(nombre_archivo, contenido):
    """Devuelve una lista de dicts {id, nombre, apellido} (id puede ser None)."""
    extension = nombre_archivo.rsplit(".", 1)[-1].lower() if "." in nombre_archivo else ""
    if extension == "xlsx":
        filas = _filas_xlsx(contenido)
    elif extension in ("csv", "txt"):
        filas = _filas_csv(contenido)
    elif extension == "xls":
        raise ErrorImportacion("Los archivos .xls (Excel viejo) no se pueden leer: guardalo como .xlsx o .csv.")
    else:
        raise ErrorImportacion("El archivo tiene que ser una planilla .csv o .xlsx.")

    # Se saltean las filas vacías del principio.
    while filas and not any(_texto(v) for v in filas[0]):
        filas.pop(0)
    if not filas:
        raise ErrorImportacion("El archivo está vacío.")

    posiciones = _columnas(filas[0])
    participantes, errores, ids = [], [], set()
    for numero, fila in enumerate(filas[1:], start=2):
        def celda(campo):
            i = posiciones.get(campo)
            return _texto(fila[i]) if i is not None and i < len(fila) else ""

        nombre, apellido, id_texto = celda("nombre"), celda("apellido"), celda("id")
        if not (nombre or apellido or id_texto):
            continue  # fila vacía
        if not nombre or not apellido:
            errores.append(f"fila {numero}: falta el nombre o el apellido")
            continue
        if len(nombre) > LARGO_MAXIMO or len(apellido) > LARGO_MAXIMO:
            errores.append(f"fila {numero}: el nombre o el apellido es demasiado largo")
            continue
        participante_id = None
        if id_texto:
            if not re.fullmatch(r"\d{1,9}", id_texto) or int(id_texto) == 0:
                errores.append(f"fila {numero}: el id «{id_texto}» no es un número entero positivo")
                continue
            participante_id = int(id_texto)
            if participante_id in ids:
                errores.append(f"fila {numero}: el id {participante_id} está repetido")
                continue
            ids.add(participante_id)
        participantes.append({"id": participante_id, "nombre": nombre, "apellido": apellido})

    if errores:
        detalle = "; ".join(errores[:5]) + (f" (y {len(errores) - 5} más)" if len(errores) > 5 else "")
        raise ErrorImportacion(f"No se importó nada porque hay errores en el archivo: {detalle}.")
    if not participantes:
        raise ErrorImportacion("El archivo no tiene participantes (solo la fila de títulos).")
    if len(participantes) > MAX_FILAS:
        raise ErrorImportacion(f"El archivo tiene más de {MAX_FILAS} participantes.")
    return participantes
